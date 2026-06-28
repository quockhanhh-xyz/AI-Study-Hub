from __future__ import annotations

import json
import os
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn


SOURCE = Path(r"C:\tmp\detail-mainflow.docx")
OUT_DIR = Path(r"D:\SWP391\AI-Study-Hub\.codex-tmp\detail-mainflow-extracted")


def iter_blocks(parent):
    body = parent.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def paragraph_images(paragraph, document):
    found = []
    for blip in paragraph._p.xpath(".//a:blip"):
        rel_id = blip.get(qn("r:embed"))
        if not rel_id:
            continue
        part = document.part.related_parts.get(rel_id)
        if part is not None:
            found.append({"rel_id": rel_id, "partname": str(part.partname)})
    return found


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document(SOURCE)
    blocks = []
    for index, block in enumerate(iter_blocks(doc), start=1):
        if isinstance(block, Paragraph):
            blocks.append(
                {
                    "index": index,
                    "type": "paragraph",
                    "style": block.style.name if block.style else None,
                    "text": block.text,
                    "images": paragraph_images(block, doc),
                }
            )
        else:
            rows = []
            for row in block.rows:
                rows.append([cell.text for cell in row.cells])
            blocks.append({"index": index, "type": "table", "rows": rows})

    media_dir = OUT_DIR / "media"
    media_dir.mkdir(exist_ok=True)
    with ZipFile(SOURCE) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        for name in media:
            target = media_dir / Path(name).name
            target.write_bytes(archive.read(name))

    result = {
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "inline_shape_count": len(doc.inline_shapes),
        "media_files": sorted(os.listdir(media_dir)),
        "blocks": blocks,
    }
    (OUT_DIR / "content.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps({key: result[key] for key in result if key != "blocks"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
